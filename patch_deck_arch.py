"""
Replaces slide 5 of full_research_deck.pptx with the architecture diagram image
+ a compact equation/variant panel below it.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DARK  = RGBColor(0x12,0x14,0x1A); DARK2 = RGBColor(0x10,0x14,0x20)
CARD  = RGBColor(0x1E,0x22,0x2C); BORDER= RGBColor(0x35,0x3C,0x50)
BLUE  = RGBColor(0x45,0x9B,0xFF); TEAL  = RGBColor(0x45,0xC4,0xFF)
GREEN = RGBColor(0x2E,0xCC,0x71); ORNG  = RGBColor(0xFF,0x7F,0x2A)
GOLD  = RGBColor(0xFF,0xC1,0x07); RED   = RGBColor(0xFF,0x4D,0x4D)
PURP  = RGBColor(0xA0,0x6C,0xFF); MUTED = RGBColor(0x8A,0x93,0xA8)
WHITE = RGBColor(0xFF,0xFF,0xFF)

W = Inches(13.33); H = Inches(7.5)

def bx(sl, l, t, w, h, c=CARD):
    s = sl.shapes.add_shape(1, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = c
    s.line.fill.background(); return s

def bd(sl, l, t, w, h, c=CARD, bc=BLUE, bp=1.0):
    s = sl.shapes.add_shape(1, l, t, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = c
    s.line.color.rgb = bc; s.line.width = Pt(bp); return s

def tx(sl, text, l, t, w, h, sz=11, bold=False, col=WHITE,
       align=PP_ALIGN.LEFT, italic=False):
    tb = sl.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p  = tf.paragraphs[0]; p.alignment = align
    r  = p.add_run(); r.text = text
    r.font.size = Pt(sz); r.font.bold = bold
    r.font.color.rgb = col; r.font.italic = italic
    return tb

def ln(sl, x1, y1, x2, y2, c=MUTED, pt=0.75):
    cn = sl.shapes.add_connector(1, x1, y1, x2, y2)
    cn.line.color.rgb = c; cn.line.width = Pt(pt)

# ── load and clear slide 5 ────────────────────────────────────────────────
prs = Presentation("/home/lam23005/STEM-GNN/full_research_deck.pptx")
sl  = prs.slides[4]

sp_tree = sl.shapes._spTree
for sp in list(sl.shapes):
    sp_tree.remove(sp._element)

fill = sl.background.fill
fill.solid(); fill.fore_color.rgb = DARK

bx(sl, 0, 0, W, Inches(0.065), c=BLUE)
bx(sl, 0, H-Inches(0.065), W, Inches(0.065), c=BLUE)

# Header
tx(sl, "Proposed Method  —  Trust-Gated Architecture",
   Inches(0.55), Inches(0.12), Inches(12.4), Inches(0.55),
   sz=24, bold=True)
tx(sl, "Trust gate taps e_v BEFORE VQ quantization — cleaner semantic signal than STEM-GNN's post-VQ router",
   Inches(0.55), Inches(0.72), Inches(12.4), Inches(0.32),
   sz=11.5, col=MUTED)
ln(sl, Inches(0.55), Inches(1.1), Inches(12.78), Inches(1.1), c=BLUE, pt=0.9)

# Architecture diagram (main image)
sl.shapes.add_picture(
    "/home/lam23005/STEM-GNN/arch_diagram.png",
    Inches(0.4), Inches(1.18), Inches(12.55), Inches(4.12),
)

# Bottom panel — equation + variants
bd(sl, Inches(0.4), Inches(5.42), Inches(12.55), Inches(1.93),
   c=DARK2, bc=BORDER, bp=0.8)

# Equation
tx(sl, "h_v  =  β_v · W_self(e_v)  +  (1 − β_v) · W_neigh( MEAN{ e_u : u ∈ N(v) } )",
   Inches(0.6), Inches(5.52), Inches(12.2), Inches(0.48),
   sz=14.5, bold=True, col=TEAL, align=PP_ALIGN.CENTER)

ln(sl, Inches(0.6), Inches(6.05), Inches(12.78), Inches(6.05), c=BORDER, pt=0.6)

# Gate variants in a row
variants = [
    ("A  Cosine",      "1 − σ(cos(h_v, mean h_u))",           RED),
    ("B  Confidence",  "1 − mean[max softmax(logits_u)]",      ORNG),
    ("C  Learned",     "MLP(e_v ‖ mean e_u).sigmoid()",        PURP),
    ("D  LLM cosine",  "1 − σ(cos(e_v, mean e_u))",           TEAL),
    ("Proposed",       "LLM('is u relevant for v?') → β_v",   GREEN),
]
for vi, (name, eq, col) in enumerate(variants):
    vl = Inches(0.55) + vi * Inches(2.54)
    bx(sl, vl, Inches(6.12), Inches(0.06), Inches(1.15),
       c=col)
    tx(sl, name, vl + Inches(0.12), Inches(6.12), Inches(2.38), Inches(0.34),
       sz=10, bold=True, col=col)
    tx(sl, eq, vl + Inches(0.12), Inches(6.48), Inches(2.38), Inches(0.82),
       sz=8.8, col=WHITE, italic=True)

prs.save("/home/lam23005/STEM-GNN/full_research_deck.pptx")
print("Slide 5 updated with architecture diagram.")
